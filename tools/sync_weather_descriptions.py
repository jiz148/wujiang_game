"""Copy source weather rules into hero descriptions without changing skill parsing.

The original skill/trait columns are intentionally left intact: their fragments
are used by the battle implementation. Run this script after changing the
weather sheet or a hero's weather references.
"""

from __future__ import annotations

import argparse
from pathlib import Path
from pprint import pformat

from openpyxl import load_workbook
from openpyxl.styles import Alignment


ROOT = Path(__file__).resolve().parents[1]
WORKBOOK = ROOT / "data" / "武将yoo.xlsx"
OUTPUT = ROOT / "src" / "wujiang" / "tactical" / "heroes" / "weather_effects_data.py"
DOC_OUTPUT = ROOT / "docs" / "武将天气效果清单.md"
HEADER = "天气效果说明"

# These descriptions incorporate the detailed rules already settled in the
# gameplay documentation; the remaining entries come directly from the source
# workbook's 天气 sheet.
IMPLEMENTED_RULES = {
    "沙尘": "各单位自己的回合结束时，土属性单位及其召唤物不受伤；飞行单位失去1/8生命，其他单位失去1/16生命。范围内不能隐身，已有隐身失效，回避距离-1。同名来源不叠加。",
    "大雪崩": "双方非冰属性武将及其召唤物不能移动或使用主动技能；冰属性武将及其召唤物速+2；“雪崩”技能伤害值+1且破魔。",
    "万魔殿": "双方恶魔武将及原召唤链攻+1；每次完成普攻独立以1/2概率增加本全局回合1次普攻。集团恶魔首领自身另有速+3特性。",
    "天空圣域": "双方天使武将及原召唤链魔上限+1，每个自己的行动槽开始时血+1/4、当前魔+1，并免疫破魔或半破魔攻击、技能的对应伤害和附效片段；与“天空的圣域”同名。",
    "湿地草原": "双方当前等级1的真正武将本体普攻和反击破魔；不包括分身、召唤物或军兵，技能不因此破魔。",
    "无常之雾": "除无常、侯鸟标记保护及效果免疫单位外，单位正式声明普攻或主动技能时独立以1/2概率失败；失败仍消耗费用、次数和冷却。",
    "王者的看破": "非同队飞王实际声明被动或反应技能时固定失去3/4生命，随后在场同队飞王本体各回复2魔；持续到当前全局回合结束。",
}

SOURCE_RULE_CLARIFICATIONS = {
    "大海": "水属性武将及其召唤物普攻伤害+1、速+1、技能威力+1，并获得穿人；其他单位伤害性技能威力-1（原表写作“伤害zai1以上”，具体阈值待明确）。",
    "寒霜": "冰属性武将及其召唤物以外的单位速-1；使用技能时“魔+1”（原表未说明是费用增加还是当前魔增加，待明确）。",
}

EXTRA_RULES = {
    "毒蛇沼泽": "原始资料只写明持续到下个回合结束前，未定义这种天气的独立效果；尚未实装。",
    "天神的狂沙": "持续2轮；场上单位只能使用阿奴比斯的黄金柜已宣告的技能，其他技能不能使用；尚未实装。",
    "禁术。大寒波": "所选2×6区域内，没有冰帝随从。里次元窃习者的一方不能攻击或使用技能，持续3轮；尚未实装。",
    "电磁场": "所有单位使用魔时失去1/8生命；尚未实装。",
}

ROW_EXTRA_NAMES = {
    27: ("无常之雾",),
    28: ("王者的看破",),
    188: ("天空圣域",),
    286: ("毒蛇沼泽",),
    288: ("天神的狂沙",),
    298: ("禁术。大寒波",),
    305: ("电磁场",),
}


def collect(*, check: bool = False) -> tuple[dict[int, str], dict[int, str]]:
    workbook = load_workbook(WORKBOOK)
    heroes = workbook["最新武将"]
    weather = workbook["天气"]
    rules = {
        str(weather.cell(row, 1).value).strip(): str(weather.cell(row, 2).value).strip()
        for row in range(1, weather.max_row + 1)
        if weather.cell(row, 1).value and weather.cell(row, 2).value
    }
    rules.pop("ex:", None)
    rules.update(SOURCE_RULE_CLARIFICATIONS)
    rules.update(IMPLEMENTED_RULES)
    rules.update(EXTRA_RULES)
    by_row: dict[int, str] = {}
    by_name: dict[int, str] = {}
    for row in range(2, heroes.max_row + 1):
        name = heroes.cell(row, 5).value
        if not name:
            continue
        raw = f"{heroes.cell(row, 11).value or ''} {heroes.cell(row, 12).value or ''}"
        names = {weather_name for weather_name in rules if weather_name in raw}
        names.update(ROW_EXTRA_NAMES.get(row, ()))
        if row == 380:  # “荆棘诅咒” is a skill name, not the 荆棘 weather.
            names.discard("荆棘")
        if row == 67:  # This trait refers to all beneficial weather effects.
            names.add("通用天气")
        if not names and "天气" not in raw:
            continue
        if not names:
            raise ValueError(f"Weather reference without a rule at source row {row}: {name}")
        parts = []
        for weather_name in sorted(names):
            if weather_name == "通用天气":
                parts.append("通用天气：可应用实际生效天气的有利效果，具体效果按对应天气规则判断；该特性尚未实装。")
            else:
                parts.append(f"{weather_name}：{rules[weather_name]}")
        by_row[row] = "\n".join(parts)
        by_name[row] = str(name)
        if check:
            if heroes.cell(row, 13).value != by_row[row]:
                raise ValueError(f"Weather description out of date at source row {row}: {name}")
        else:
            cell = heroes.cell(row, 13)
            cell.value = by_row[row]
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    if heroes.cell(1, 13).value not in (None, HEADER):
        raise ValueError("Column M is already in use")
    if check and heroes.cell(1, 13).value != HEADER:
        raise ValueError("Weather description header is missing")
    if check:
        for row in range(2, heroes.max_row + 1):
            if row not in by_row and heroes.cell(row, 13).value:
                raise ValueError(f"Stale weather description at source row {row}")
    if not check:
        heroes.cell(1, 13).value = HEADER
        heroes.column_dimensions["M"].width = 72
        workbook.save(WORKBOOK)
    return by_row, by_name


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--check", action="store_true", help="Check generated descriptions without writing")
    args = parser.parse_args()
    by_row, by_name = collect(check=args.check)
    workbook = load_workbook(WORKBOOK, data_only=True)
    weather_sheet = workbook["天气"]
    weather_rows = [
        (row, str(weather_sheet.cell(row, 1).value).strip(), str(weather_sheet.cell(row, 2).value).strip())
        for row in range(1, weather_sheet.max_row + 1)
        if weather_sheet.cell(row, 1).value and weather_sheet.cell(row, 2).value
    ]
    weather_rows = [item for item in weather_rows if item[1] != "ex:"]
    body = (
        "# Generated from data/武将yoo.xlsx column M by tools/sync_weather_descriptions.py.\n"
        "# Do not edit by hand.\n\n"
        "WEATHER_HERO_EFFECTS_BY_ROW = " + pformat(by_row, width=110, sort_dicts=True) + "\n\n"
        "WEATHER_HERO_EFFECTS_BY_NAME = " + pformat({by_name[row]: text for row, text in by_row.items() if row < 20}, width=110, sort_dicts=True) + "\n"
    )
    doc_lines = [
        "# 武将天气效果与实现状态",
        "",
        "本清单按 `data/武将yoo.xlsx` 的“天气”表和武将技能、特性生成。原表技能和特性文字保持原样；“最新武将”表的“天气效果说明”列逐位补充具体效果。游戏内选将悬浮信息、房间详情和战斗武将信息读取同一份补充说明。",
        "",
        "## 已实现的天气",
        "",
    ]
    for name in ("沙尘", "大雪崩", "万魔殿", "天空圣域", "湿地草原", "无常之雾", "王者的看破"):
        doc_lines.append(f"- **{name}**：{IMPLEMENTED_RULES[name]}")
    doc_lines.extend(["", "沙尘支持全场和随岩神移动的局部区域；天空圣域支持全场和随制裁者移动的局部区域。常规全场天气双方各可保有一份，跨来源同名效果不叠加。无常之雾和王者的看破是按各自规则追加的专属天气，不占原始天气表的编号。", "", "## 原始天气表中尚未实现的天气", ""])
    for row, name, rule in weather_rows:
        if name not in IMPLEMENTED_RULES:
            doc_lines.append(f"- **{name}**（天气表第 {row} 行，设计规则，战斗中尚未实现）：{rule}")
    doc_lines.extend(["", "## 原始天气表之外、尚未实现的天气", ""])
    for name, rule in EXTRA_RULES.items():
        doc_lines.append(f"- **{name}**：{rule}")
    doc_lines.extend(["", "## 涉及天气的武将说明", "", "下面每一项也写入原始武将表的“天气效果说明”列。", ""])
    for row, description in by_row.items():
        doc_lines.append(f"- **{by_name[row]}**（源表第 {row} 行）：{description.replace(chr(10), '；')}" )
    doc_lines.extend(["", "## 源资料待明确之处", "", "- 毒蛇沼泽只写了持续时间，没有独立天气效果定义；不能据此实现额外伤害或减速。", "- 浮鹤城的武将技能写的是主动技能有一半概率成功，而天气表写的是其他单位不能用主动技能；目前未实现，需先统一规则。", "- 大海天气表有“伤害zai1以上”的原文笔误；相关减威力阈值目前只能按原文保留，需明确后再实现。", ""])
    doc_body = "\n".join(doc_lines)
    if args.check:
        if not OUTPUT.exists() or OUTPUT.read_text(encoding="utf-8") != body:
            raise SystemExit("Weather descriptions are out of sync; run tools/sync_weather_descriptions.py")
        if not DOC_OUTPUT.exists() or DOC_OUTPUT.read_text(encoding="utf-8") != doc_body:
            raise SystemExit("Weather document is out of sync; run tools/sync_weather_descriptions.py")
    else:
        OUTPUT.write_text(body, encoding="utf-8")
        DOC_OUTPUT.write_text(doc_body, encoding="utf-8")
    print(f"{len(by_row)} weather-related heroes checked")


if __name__ == "__main__":
    main()
